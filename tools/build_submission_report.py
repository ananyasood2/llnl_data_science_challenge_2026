from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "final_submission"
OUT.mkdir(parents=True, exist_ok=True)

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
GRAY = RGBColor(80, 80, 80)

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), color); tcPr.append(shd)

def set_cell(cell, text, bold=False, color=None):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(text)); r.bold = bold; r.font.size = Pt(9); r.font.name = "Calibri"; r._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    if color: r.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        set_cell(t.rows[0].cells[i], h, True, RGBColor(0,0,0)); shade(t.rows[0].cells[i], "E8EEF5")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row): set_cell(cells[i], value)
    if widths:
        for row in t.rows:
            for cell, width in zip(row.cells, widths): cell.width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

def add_caption(doc, text):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(9)
    r = p.add_run(text); r.italic = True; r.font.size = Pt(9); r.font.color.rgb = GRAY

def add_figure(doc, relative, caption, width=6.15):
    path = ROOT / relative
    if path.exists():
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(path), width=Inches(width))
        add_caption(doc, caption)

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}"); p.add_run(text)

def para(doc, text, bold_lead=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6); p.paragraph_format.line_spacing = 1.1
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True; p.add_run(text[len(bold_lead):])
    else: p.add_run(text)

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet"); p.paragraph_format.space_after = Pt(3); p.add_run(text)

def page_break(doc): doc.add_page_break()

doc = Document()
sec = doc.sections[0]
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(0.85)
sec.header_distance = sec.footer_distance = Inches(0.45)

styles = doc.styles
normal = styles["Normal"]; normal.font.name = "Calibri"; normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri"); normal.font.size = Pt(10.5)
for name, size, color in [("Heading 1", 16, BLUE), ("Heading 2", 13, BLUE), ("Heading 3", 11.5, DARK)]:
    s=styles[name]; s.font.name="Calibri"; s._element.rPr.rFonts.set(qn("w:ascii"),"Calibri"); s.font.size=Pt(size); s.font.color.rgb=color
    s.paragraph_format.space_before=Pt(12 if name=="Heading 1" else 8); s.paragraph_format.space_after=Pt(5)

header = sec.header.paragraphs[0]; header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
rr = header.add_run("LLNL DSSI 2026 | Lattice CT Inspection Submission"); rr.font.size = Pt(8); rr.font.color.rgb = GRAY
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run("Technical report | Registered CT/graph workflow | July 2026"); fr.font.size = Pt(8); fr.font.color.rgb = GRAY

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(80); p.paragraph_format.space_after=Pt(8)
r=p.add_run("Lattice CT Inspection\nFinal Technical Submission"); r.bold=True; r.font.size=Pt(25); r.font.color.rgb=DARK
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(30)
r=p.add_run("Part 1: Unit-cell segmentation and NDE\nPart 2: Registered 0.5%-missing octet-lattice screening"); r.font.size=Pt(14); r.font.color.rgb=GRAY
table(doc,["Submission scope","Status"],[
    ["Part 1", "Completed: threshold optimization, segmentation, skeleton, NDE visual report"],
    ["Part 2", "Completed: registered graph/CT screening and sensitivity analysis"],
    ["Quantitative defect confirmation", "Not supported by available registration evidence; no missing struts confirmed"],
], [2.1,4.3])
para(doc,"Prepared from the repository's saved, traceable analysis artifacts. The unregistered STL files were excluded from all quantitative CT comparisons.")
page_break(doc)

heading(doc,"Executive Summary")
para(doc,"This submission consolidates the completed Part 1 and Part 2 outputs for the LLNL Data Science Challenge. Part 1 establishes a reproducible segmentation-to-NDE workflow on the unit-cell CT volume. Part 2 evaluates the supplied registered 0.5%-missing octet-lattice CT/JSON pair using centerline tube support and sensitivity checks.")
table(doc,["Finding","Result","Interpretation"],[
    ["Part 1 selected threshold","0.005813","Otsu-selected balance of strut continuity and mask compactness."],
    ["Part 1 mask","717,852 voxels (4.2787%)","Selected mask is shape-compatible with the original 256 x 256 x 256 volume."],
    ["Part 1 skeleton","3,182 voxels; 1 component","Connected unit-cell structural proxy."],
    ["Part 2 design graph","18,468 struts; 10,206 junctions","Supplied registered JSON; used without an STL comparison."],
    ["Part 2 retained operating point","Tube radius 2 voxels; threshold 39,725","Least-flagging candidate setting of the tested radii."],
    ["Part 2 confirmation result","0 confirmed missing struts","Candidate behavior is not stable enough for a physical-defect claim."],
], [1.75,1.65,3.0])
para(doc,"Decision: The workflow is complete and auditable. The Part 2 candidate tables are suitable for targeted human review, but they must not be presented as a measured missing-strut count or forced to the nominal approximately 92 removals.", "Decision: ")

heading(doc,"1. Objectives and Data")
para(doc,"The analysis has two linked objectives: (1) construct and document a segmentation, skeletonization, and NDE-report workflow for a unit-cell CT dataset; and (2) screen an aligned octet-lattice specimen for missing or disconnected strut candidates using only the supplied registered CT graph.")
table(doc,["Dataset","Input","Use"],[
    ["Part 1 unit cell","data/unitcell/unitcell.npy","Intensity volume for segmentation and NDE."],
    ["Part 2 CT","0.5%-missing TIFF stack, 761 x 815 x 837 uint16","As-built image evidence."],
    ["Part 2 design","Registered JSON, 18,468 struts","Expected centerlines in CT coordinates."],
    ["STL variants","Unregistered","Excluded from quantitative comparison."],
], [1.25,2.6,2.55])

heading(doc,"2. Part 1 Methods and Results")
para(doc,"Three segmentation thresholds were compared. The selected threshold, 0.005813, is the Otsu value and provided a balanced foreground fraction with visually continuous lattice material. The chosen binary mask was skeletonized and rendered from two 3D perspectives for NDE review.")
table(doc,["Threshold","Foreground fraction","Decision"],[
    ["0.004500","4.3416%","Lower threshold candidate."],
    ["0.005813","4.2787%","Selected Otsu threshold."],
    ["0.007000","4.2480%","Higher threshold candidate."],
], [1.6,1.8,3.0])
add_figure(doc,"output/part1/unitcell_raw_slice_128.png","Figure 1. Part 1 unit-cell raw CT slice (slice 128).",4.8)
add_figure(doc,"output/part1/unitcell_mask_slice_128.png","Figure 2. Part 1 selected segmentation mask at the same slice.",4.8)
add_figure(doc,"output/part1/nde_report/view_a.png","Figure 3. Part 1 NDE View A: elevation 30 degrees, azimuth 45 degrees.",5.2)
add_figure(doc,"output/part1/nde_report/view_b.png","Figure 4. Part 1 NDE View B: elevation 60 degrees, azimuth 45 degrees.",5.2)
para(doc,"Part 1 conclusion: the selected mask and skeleton are dimensionally compatible with the source volume. The single connected skeletal component supports the intended unit-cell connectivity for this dataset; it is a morphology proxy, not a mechanical simulation.")

page_break(doc)
heading(doc,"3. Part 2 Registered CT-to-Graph Screening")
para(doc,"The Part 2 graph was already registered to the CT voxel coordinate system. Expected strut centerlines were sampled after converting graph XYZ coordinates to CT indexing. A local tube test used the 75th-percentile CT intensity around each sample, which is more tolerant of small offsets than a single centerline voxel. Samples were trimmed to the central 8%-92% of each strut to reduce junction influence.")
bullet(doc,"31 samples per expected strut; seven representative CT slices for threshold candidates.")
bullet(doc,"Candidate labels: low-support (missing candidate), internal-gap (disconnected candidate), intermediate-support (uncertain candidate), and clear.")
bullet(doc,"No unregistered STL geometry was opened or compared to the CT result.")
table(doc,["Parameter","Retained value","Reason"],[
    ["Tube radius","2 original CT voxels","Least flagging among radii 1, 2, and 3."],
    ["Selected threshold","39,725","Separation criterion while retaining median centerline support >= 0.60."],
    ["Missing rule","Material fraction < 0.04","Conservative low-support candidate screen."],
    ["Disconnected rule","Material fraction < 0.35 with long internal gap","Separates gap-like screening behavior from low support."],
], [1.45,1.65,3.3])
add_figure(doc,"output/part2/refined_registration_20260728/tube_r2/threshold_diagnostics.png","Figure 5. Part 2 radius-2 centerline intensity distribution and threshold-candidate separation.",6.15)

heading(doc,"4. Part 2 Sensitivity and Candidate Results")
table(doc,["Tube radius (voxels)","Threshold","Low support","Disconnected","Uncertain","Clear"],[
    ["1","39,725","1,217","5,614","1,199","10,438"],
    ["2 (retained)","39,725","1,049","5,262","1,083","11,074"],
    ["3","42,017","1,159","6,505","1,384","9,420"],
], [1.1,1.0,1.05,1.1,1.0,1.1])
para(doc,"At the retained radius-2 operating point, 7,394 struts are review candidates: 1,049 low-support, 5,262 disconnected, and 1,083 uncertain. This is a screening result, not a defect count. The marked dependence on tube radius and threshold is stronger evidence for residual registration, segmentation, and partial-volume effects than for thousands of physical defects.")
add_figure(doc,"output/part2/registered_screen_20260728/ct_mask_evidence_overlay.png","Figure 6. Supplemental CT/segmentation overlay with nearby image-derived flags. This is supporting visual evidence only; it is not the authoritative result source.",6.15)

heading(doc,"5. Discussion and Validation Status")
para(doc,"The nominal 0.5% label implies approximately 92 intentionally removed struts out of 18,468, but this value was never used as a target during classification. The screen returns far more candidates than the nominal label and does so inconsistently as parameters change. A valid analysis must retain that discrepancy rather than tune the output to meet the nominal expectation.")
para(doc,"The supplied JSON records registered graph positions but does not provide a transform covariance, independent fiducials, or released per-strut ground truth. Therefore an additional rigid or deformable registration refinement cannot be defensibly estimated and persisted from the currently available artifacts. The appropriate conclusion is no confirmed physical missing struts, with a ranked candidate list for review.")
table(doc,["Claim","Status","Evidence"],[
    ["Registered JSON and CT were used","Supported","Paths and dimensions stored in each inspection summary."],
    ["Unregistered STL was excluded","Supported","Workflow and artifact provenance state this explicitly."],
    ["Candidates are reproducible","Supported","Per-strut CSV and threshold diagnostics are saved."],
    ["Physical missing-strut count is known","Not supported","No registered truth join or blinded neighborhood adjudication."],
    ["Mechanical repair benefit is known","Not supported","Historical repair outputs are non-FEA candidate scenarios only."],
], [2.25,1.15,2.85])

heading(doc,"6. Recommended Next Steps")
bullet(doc,"Obtain independent registration landmarks or transform uncertainty information for the CT/graph pair.")
bullet(doc,"Perform blinded CT-neighborhood adjudication of ranked candidates, then record reviewer decisions.")
bullet(doc,"If registered per-strut truth becomes available, calculate precision, recall, F1, and a confusion matrix before reporting a defect rate.")
bullet(doc,"Use the saved radius-1/radius-2/radius-3 outputs as a sensitivity baseline for any future improved registration.")

page_break(doc)
heading(doc,"Appendix A. Deliverable Manifest")
para(doc,"The following files are the complete submission-relevant output set. Large arrays and the 18,468-row candidate table are provided as companion artifacts rather than embedded in this PDF/DOCX.")
table(doc,["Part","Artifact group","Contents"],[
    ["1","threshold_optimizer/","Three masks, three slice comparisons, and threshold selection memo."],
    ["1","unitcell_mask.npy; unitcell_skeleton.npy","Selected segmentation and skeleton arrays."],
    ["1","nde_report/","Markdown NDE report and two 3D renders."],
    ["1","registered_strut_screen/","Registered strut screen CSV and JSON."],
    ["2","metadata/; graph_comparison/","Dataset inventory and registered graph comparison artifacts."],
    ["2","refined_registration_20260728/","Authoritative readout plus radius-1, radius-2, and radius-3 runs."],
    ["2","tube_r2/strut_candidates.csv","All 18,468 per-strut candidate records for the retained setting."],
    ["2","registered_screen_20260728/","Supplemental distance-screen sensitivity, overlay, provenance, and interactive scene."],
    ["2","visual_repair*/","Historical non-FEA candidate-repair scenarios; not physical recommendations."],
], [0.55,2.35,3.35])

heading(doc,"Appendix B. Reproducibility and Audit Notes")
bullet(doc,"Part 1 source volume: data/unitcell/unitcell.npy. Selected threshold: 0.005813.")
bullet(doc,"Part 2 source CT: data/missing_struts/tif_stacks/210127_Brian_Tran_strut_lattices_0point5dash1 1 Slices.tif.")
bullet(doc,"Part 2 registered graph: data/missing_struts/registered_jsons/210127_Brian_Tran_strut_lattices_0point5dash1 1 Slices.json.")
bullet(doc,"Authoritative Part 2 output: output/part2/refined_registration_20260728/tube_r2/.")
bullet(doc,"Full artifact listing: output/DELIVERABLES.md.")
para(doc,"End of submission report.")

doc.save(OUT / "Lattice_CT_Inspection_Final_Submission.docx")
print(OUT / "Lattice_CT_Inspection_Final_Submission.docx")
